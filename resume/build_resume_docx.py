from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT = Path(__file__).resolve().parent / "郑书振-前端开发工程师-简历.docx"
FONT = "微软雅黑"
INK = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x6B, 0x72, 0x80)
ACCENT = RGBColor(0x31, 0x3A, 0x8F)
LINE = "D9DDF0"


def set_run(run, size=9.5, bold=False, color=INK, italic=False):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    return run


def para(doc, before=0, after=2, line=1.12, keep=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    p.paragraph_format.keep_with_next = keep
    return p


def bottom_border(p, color=LINE, size="6"):
    pPr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "3")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)
    pPr.append(borders)


def section(doc, title):
    p = para(doc, before=10, after=4, keep=True)
    set_run(p.add_run(title), size=12, bold=True, color=ACCENT)
    bottom_border(p)


def entry(doc, left, right):
    p = para(doc, before=4, after=1.5, keep=True)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(17.2), WD_TAB_ALIGNMENT.RIGHT)
    set_run(p.add_run(left), size=10, bold=True)
    p.add_run("\t")
    set_run(p.add_run(right), size=9, color=MUTED)


def bullet(doc, text, bold_prefix=None):
    p = para(doc, before=0, after=1.5)
    p.paragraph_format.left_indent = Cm(0.45)
    p.paragraph_format.first_line_indent = Cm(-0.45)
    set_run(p.add_run("•  "), size=9.5, color=ACCENT, bold=True)
    if bold_prefix:
        set_run(p.add_run(bold_prefix), size=9.5, bold=True)
    set_run(p.add_run(text), size=9.5)
    return p


def tech(doc, text):
    p = para(doc, before=0, after=2)
    p.paragraph_format.left_indent = Cm(0.45)
    set_run(p.add_run("技术栈："), size=8.7, bold=True, color=MUTED)
    set_run(p.add_run(text), size=8.7, color=MUTED)


def main():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(1.35)
    sec.bottom_margin = Cm(1.25)
    sec.left_margin = Cm(1.9)
    sec.right_margin = Cm(1.9)

    # 默认样式
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = INK
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.space_after = Pt(2)

    # 姓名与抬头
    p = para(doc, before=0, after=1)
    set_run(p.add_run("郑书振"), size=23, bold=True)
    p = para(doc, before=0, after=2)
    set_run(p.add_run("前端开发工程师"), size=11, bold=True, color=ACCENT)
    set_run(p.add_run("  ·  杭州  ·  4 年前端开发经验"), size=10, color=MUTED)
    p = para(doc, before=0, after=1)
    set_run(p.add_run("[手机号]  ·  [邮箱]  ·  [GitHub/博客]"), size=9, color=MUTED)
    p = para(doc, before=0, after=5)
    set_run(p.add_run("求职意向：前端开发工程师  ｜  期望薪资：14–18K  ｜  随时到岗"), size=9, color=MUTED)
    bottom_border(p, color="A5B4FC", size="8")

    section(doc, "个人优势")
    bullet(doc, "4 年前端开发经验，聚焦电商商家端、培训考试、桌面端，能独立负责需求分析、技术方案、开发联调、上线与运维。")
    bullet(doc, "熟悉 Vue2/3、React、TypeScript、Electron；有 NestJS + PostgreSQL 全栈项目经验，能独立完成管理端/BFF 接口。")
    bullet(doc, "有复杂业务落地经验：活动创编/报名链路、题库考试、人脸识别、蓝牙硬件通信、WebSocket、数据大屏、PDF 报表。")
    bullet(doc, "注重工程化与可维护性，熟悉 Vite/Webpack、pnpm、Git、Nginx，有性能优化与线上问题排查经验。")

    section(doc, "技能栈")
    bullet(doc, "HTML/CSS/JavaScript/TypeScript、Vue2/3、React、Pinia、Vue Router、Element Plus、Ant Design、Vant", "前端：")
    bullet(doc, "uni-app、微信小程序、Electron", "跨端/桌面：")
    bullet(doc, "ECharts、数据大屏、tracking.js、face-api.js、html2canvas、jsPDF", "可视化/交互：")
    bullet(doc, "Vite、Webpack、pnpm、Git、Nginx", "工程化/部署：")
    bullet(doc, "NestJS、PostgreSQL、MySQL、Redis（了解）、Prisma", "后端/数据：")

    section(doc, "工作经历")
    entry(doc, "浙江外企德科人力资源服务有限公司（阿里淘天集团外包）", "2025.12–2026.09")
    p = para(doc, before=0, after=2)
    set_run(p.add_run("前端开发工程师"), size=9.3, color=MUTED)
    bullet(doc, "驻场阿里妈妈/淘天商家端，负责「品牌新享」「智能托管」等商家营销产品的 PC/千牛端开发。")
    bullet(doc, "负责品牌新享活动创编与报名链路：按新客、老客、新品、会员加速等场景选品，配置礼金/费率、报名活动、签署协议并提交。")
    bullet(doc, "负责千牛端品牌新享/智能托管页面开发与调试，参与顶部 banner、数据报表、老活动链路升级。")
    bullet(doc, "与产品、后端、测试协作，完成需求评审、联调、测试环境部署与上线，独立负责前端项目交付。")
    tech(doc, "Magix、JavaScript、HTML/CSS、千牛端")

    entry(doc, "杭州广域软件有限公司", "2023.03–2025.11")
    p = para(doc, before=0, after=2)
    set_run(p.add_run("前端开发工程师"), size=9.3, color=MUTED)
    bullet(doc, "独立负责展厅培训数字化产品前端，覆盖 Web 管理端、数据大屏、Electron 桌面端 3 类终端。")
    bullet(doc, "负责培训智能管控、综合测评/心理测评、自助绑卡终端、游戏活动页等项目的前端开发与交付。")
    bullet(doc, "负责需求调研、技术方案设计、测试环境部署、生产上线与后续运维，编写用户操作手册。")
    tech(doc, "React、Vue2/3、Electron、ECharts、face-api.js、WebSocket、html2canvas、jsPDF")

    entry(doc, "杭州掌派科技有限公司", "2022.01–2022.05")
    p = para(doc, before=0, after=2)
    set_run(p.add_run("前端开发工程师"), size=9.3, color=MUTED)
    bullet(doc, "参与游戏活动页、萌娘百科编辑中心开发，负责签到/抽奖、活动预约、内容编辑与文件上传。")
    tech(doc, "Vue2/Vue3、Element UI、lucky-canvas、Axios、Pinia")

    section(doc, "项目经历")
    entry(doc, "淘宝商家端「品牌新享」｜前端开发", "2025.12–2026.09")
    bullet(doc, "淘宝/天猫商家营销产品，覆盖新客、老客、新品、会员加速等场景；商家在千牛/PC 端选品、配置礼金/费率、报名活动、签署协议。", "项目描述：")
    bullet(doc, "负责活动创编/报名链路的前端开发：场景与商品选择、礼金/费率配置、报名提交、协议签署、提交校验。")
    bullet(doc, "负责千牛端品牌新享/智能托管页面开发与调试，参与 banner、数据报表、老活动链路升级。")
    bullet(doc, "与产品、后端、测试协作，完成需求评审、联调、测试环境部署与上线。")
    tech(doc, "Magix、JavaScript、HTML/CSS、千牛端")

    entry(doc, "培训人员智能管控系统｜独立负责前端", "2023–2025")
    bullet(doc, "展厅培训场景一体化管控系统，含学员管理、数据大屏、环境监测、人脸识别。", "项目描述：")
    bullet(doc, "使用 React 全家桶搭建前端架构，完成系统配置、成绩查询、数据展示等模块。")
    bullet(doc, "集成 face-api.js、tracking.js，实现人脸识别登录与身份匹配，支持学号、人脸双查询。")
    bullet(doc, "基于 ECharts 构建数据大屏与展厅热力图，通过蓝牙/感应设备实时展示人员分布；对接温湿度、空气质量设备。")
    bullet(doc, "使用 html2canvas + jsPDF 导出成绩报表；支持角色权限、培训清单、成绩统计与导出。")
    tech(doc, "React、React Router、ECharts、Ant Design、tracking.js、face-api.js、html2canvas、jsPDF")

    entry(doc, "综合测评与心理测评系统｜独立负责前端", "2023–2025")
    bullet(doc, "面向培训/测评场景的在线考试与心理评估系统，包含题库配置、考试、错题解析、心理报告。", "项目描述：")
    bullet(doc, "使用 Vue3 + Pinia 搭建系统，完成题目增删改查、分类管理，按专业配置题库、题量、分值。")
    bullet(doc, "支持视频题、图片题；测评结束后查看分值、错题详情与解析，提供复盘。")
    bullet(doc, "心理测评自动生成分析报告，支持在线查看与 PDF 下载，使用 ECharts 可视化报告数据；集成 face-api.js 人脸识别 + 学号登录双认证。")
    tech(doc, "Vue3、Pinia、Vue Router、ECharts、face-api.js、html2canvas、jsPDF")

    entry(doc, "培训人员自助绑卡终端｜独立负责前端", "2023–2025")
    bullet(doc, "展厅培训现场桌面端自助绑卡系统，学员身份识别、蓝牙手环绑定、考勤管理。", "项目描述：")
    bullet(doc, "使用 Electron + Vue3 构建跨平台桌面应用，实现前端与设备端通信；集成 face-api.js 人脸识别与学号校验。")
    bullet(doc, "通过 WebSocket 与吐卡设备实时通信，显示手环编号与设备状态；实现识别 → 感应手环 → 绑定 → 设备吐卡的完整流程。")
    bullet(doc, "手环用于人员定位/成绩采集，附考勤查询。")
    tech(doc, "Electron、Vue3、Vue Router、Element Plus、face-api.js、WebSocket")

    section(doc, "教育经历")
    entry(doc, "[学校名称]｜本科 · 软件工程", "2018.09–2022.06")

    doc.save(OUT)
    print(f"SAVED {OUT}")


if __name__ == "__main__":
    main()
